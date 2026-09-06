import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = "renewables_stocks.db"
OUTPUT_TOP100 = "top_100_leading_indian_companies_deep_esg.xlsx"

conn = sqlite3.connect(DB_PATH)
c = conn.cursor()

c.execute("""
SELECT 
    name, ticker, sector, cap_tier, market_cap_cr, close_price, pe, roce,
    re_pct, re_mw, annual_mu, grid_tariff, solar_cost, unit_shield, annual_shield_cr,
    ret_1y, ret_3y, ret_5y, ret_6y,
    brsr_status, primary_model, targets, scope1_2_reduction, plants_info
FROM companies
ORDER BY market_cap_cr DESC
LIMIT 100
""")
top100_rows = c.fetchall()

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Top 100 Leading Companies"

# Styling
font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='E2E8F0')
font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_bold = Font(name='Segoe UI', size=9, bold=True)
font_regular = Font(name='Segoe UI', size=9)

fill_navy = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
fill_header = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
fill_highlight = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
)

headers = [
    "Rank", "Company Name", "Ticker", "Sector", "Cap Tier", "Market Cap (₹ Cr)", "CMP (₹)", "P/E", "ROCE (%)",
    "RE Share (%)", "Green MW", "Annual MU", "Grid Tariff (₹)", "Solar Cost (₹)", "Unit Savings (₹)", "Annual Cost Shielded (₹ Cr)",
    "1Y Ret (%)", "3Y Ret (%)", "5Y Ret (%)", "6Y Ret (%)",
    "BRSR Status", "Sourcing Model", "Net-Zero Roadmap & Targets", "Scope 1 & 2 Reduction Progress", "Manufacturing Footprint"
]

# Header Title
ws.merge_cells("A1:Y1")
ws["A1"] = "INDIA TOP 100 LEADING CORPORATES — DEEP BRSR, CLEAN ENERGY & EQUITY PERFORMANCE"
ws["A1"].font = font_title
ws["A1"].fill = fill_navy
ws["A1"].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:Y2")
ws["A2"] = "Detailed ESG disclosures, operational captive solar/wind capacity, tariff shielding spreads, and multi-year compounded returns"
ws["A2"].font = font_subtitle
ws["A2"].fill = fill_navy
ws["A2"].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 20

# Column Headers
ws.append(headers)
header_row = ws[3]
ws.row_dimensions[3].height = 26
for cell in header_row:
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_thin

# Add Rows
for idx, r in enumerate(top100_rows, 1):
    row_data = [idx] + list(r)
    ws.append(row_data)
    current_row = ws[idx + 3]
    ws.row_dimensions[idx + 3].height = 20
    is_even = (idx % 2 == 0)
    bg_fill = fill_zebra if is_even else fill_white
    
    for c_idx, cell in enumerate(current_row, 1):
        cell.font = font_regular
        cell.fill = bg_fill
        cell.border = border_thin
        
        # Formatting
        if c_idx in [1, 3, 5]: # Rank, Ticker, Cap
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx in [2, 4]: # Name, Sector
            cell.alignment = Alignment(horizontal='left', vertical='center')
        elif c_idx in [6, 7, 16]: # MCap, CMP, Shielded Cr
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0'
            if c_idx == 16:
                cell.font = font_bold
                cell.fill = fill_highlight
        elif c_idx in [8, 9, 10, 11, 12, 13, 14, 15, 17, 18, 19, 20]: # Numbers / %
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0.0'
        else: # Text details
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[2:15]:
        val = str(cell.value or '')
        if len(val) > max_len:
            max_len = len(val)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 28
ws.column_dimensions['D'].width = 24
ws.column_dimensions['U'].width = 30
ws.column_dimensions['V'].width = 36
ws.column_dimensions['W'].width = 40
ws.column_dimensions['X'].width = 40
ws.column_dimensions['Y'].width = 40

wb.save(OUTPUT_TOP100)
print(f"Generated {OUTPUT_TOP100} with {len(top100_rows)} deep profiles successfully!")
conn.close()
